#!/usr/bin/env python3
"""The Excel file of a free competency matrix, for 1:1 meetings.

Run by hand when the matrix changes; it is not part of the site build. Commit
the output with the content change.

    python3 tools/build-matrix-xlsx.py [slug]    (default: matritsa-kompetentsij-metodista)

Needs openpyxl and PyYAML (e.g. a virtualenv: pip install openpyxl pyyaml).
Reads src/content/matrices/<slug>.md (the exported content, so no internal
notes), writes public/free/<slug>.xlsx (the path in the content's `xlsx.file`):

- «Как пользоваться» (the first sheet): the title and the lead, which column
  each step fills (and the self-check), the steps, the legend, the address
  of the page (a link with utm_source=matrix-xlsx).
- «Матрица»: one row per skill (number, category, grade, skill, ● key, *
  another role), then the marks of the methodologist and of the lead and the
  result after the 1:1 (drop-down lists «да / нет»; on * rows the result
  also «вычеркнут», as on the page: only * items can be struck out; a typed
  value outside the list is rejected), and a column for the example or the
  comment. Header row and the number column frozen, filter on, a closed row
  in light grey; printed landscape, one page wide, the header on every page.
- «Итог»: per grade the items left after striking out, closed, needed (80%,
  rounded up), key items closed, key items still open in this grade and
  below, regular items open in the grades below, and «засчитан» by the same
  rule as the page (src/lib/matrix.ts); the grade reached, or a prompt to
  fill «Итог после 1:1» while that column is empty.
- «Ключевые пункты»: grade, number, item and what to ask to see, by grade
  like the table on the page.

A step whose text starts lowercase («Отметьте сами» + «по тем же
правилам…») is joined to its title with a space, as on the page.

Black and white like the site; fonts are left to the spreadsheet app.
"""

import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SITE = 'https://danilevstratov.ru'
UTM = 'utm_source=matrix-xlsx&utm_medium=xlsx'
GRADES = [('junior', 'Junior'), ('middle', 'Middle'), ('senior', 'Senior')]
NAME = dict(GRADES)
ORDER = [g for g, _ in GRADES]

slug = sys.argv[1] if len(sys.argv) > 1 else 'matritsa-kompetentsij-metodista'
src = ROOT / 'src' / 'content' / 'matrices' / f'{slug}.md'
raw = src.read_text(encoding='utf-8')
d = yaml.safe_load(raw.split('---\n')[1])
out = ROOT / 'public' / d['xlsx']['file'].lstrip('/')
page_url = f'{SITE}/free/{slug}/'

thin = Side(style='thin', color='000000')
hair = Side(style='hair', color='999999')
bold = Font(bold=True)
wrap = Alignment(wrap_text=True, vertical='top')
grey = PatternFill('solid', fgColor='EEEEEE')


def step_line(k, s):
    # «3. Отметьте сами по тем же правилам…», «2. Попросите… . По каждому…»
    sep = ' ' if s['text'][:1].islower() else '. '
    return f"{k}. {s['title']}{sep}{s['text']}"


def header(ws):
    for c in ws[1]:
        c.font = bold
        c.alignment = wrap
        c.border = Border(bottom=thin)


def landscape(ws, title_rows=None):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if title_rows:
        ws.print_title_rows = title_rows


wb = Workbook()
hs = wb.active
hs.title = 'Как пользоваться'
ws = wb.create_sheet('Матрица')
ts = wb.create_sheet('Итог')
ks = wb.create_sheet('Ключевые пункты')

# --- Матрица -----------------------------------------------------------------
ws.append(['№', 'Категория', 'Грейд', 'Навык', 'Ключевой', 'Другая роль', 'Методист', 'Руководитель',
           'Итог после 1:1', 'Пример или комментарий'])
header(ws)
stars = []
for cat in d['categories']:
    for s in cat['skills']:
        ws.append([
            s['id'],
            f"{cat['id']}. {cat['title']}",
            NAME[s['grade']],
            s['text'],
            '●' if s.get('key') else '',
            '*' if s.get('star') else '',
            '', '', '', '',
        ])
        if s.get('star'):
            stars.append(ws.max_row)
last = ws.max_row
n = last - 1
for row in ws.iter_rows(min_row=2, max_row=last):
    for c in row:
        c.alignment = wrap
        c.border = Border(bottom=hair)
for col, w in zip('ABCDEFGHIJ', [6, 18, 9, 60, 9, 9, 11, 13, 14, 40]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'B2'
ws.auto_filter.ref = f'A1:J{last}'
landscape(ws, '1:1')


def choice(values, error):
    dv = DataValidation(type='list', formula1=f'"{",".join(values)}"', allow_blank=True,
                        showErrorMessage=True, errorTitle='Отметка', error=error)
    ws.add_data_validation(dv)
    return dv


def runs(rows):
    # [3, 4, 5, 9] -> I3:I5 I9
    out, rows = [], sorted(rows)
    for r in rows:
        if out and out[-1][1] == r - 1:
            out[-1][1] = r
        else:
            out.append([r, r])
    return [f'I{a}' if a == b else f'I{a}:I{b}' for a, b in out]


marks = choice(['да', 'нет'], 'Выберите из списка: да или нет.')
marks.add(f'G2:H{last}')
for ref in runs([r for r in range(2, last + 1) if r not in stars]):
    marks.add(ref)
struck = choice(['да', 'нет', 'вычеркнут'], 'Выберите из списка: да, нет или вычеркнут.')
for ref in runs(stars):
    struck.add(ref)
ws.conditional_formatting.add(f'A2:J{last}', FormulaRule(formula=['$I2="да"'], fill=grey))
ws.conditional_formatting.add(f'A2:J{last}', FormulaRule(formula=['AND($I2="вычеркнут",$F2="*")'], font=Font(strike=True, color='777777')))

# --- Итог --------------------------------------------------------------------
ts.append(['Грейд', 'Пунктов (без вычеркнутых)', 'Закрыто', 'Нужно (80%)', 'Ключевых', 'Ключевых закрыто',
           'Открытых ключевых в этом и ниже', 'Открытых обычных ниже', 'Засчитан'])
header(ts)
M = "'Матрица'!"
G, E, F, I = (f'{M}${c}$2:${c}${last}' for c in 'CEFI')
# «вычеркнут» takes an item out of the count only on a * row, as on the page
# (the drop-down offers it only there; this also holds for a pasted value);
# "~*" is a literal asterisk, a bare "*" in COUNTIFS matches any text
for i, (_, name) in enumerate(GRADES):
    r = i + 2
    lower = [g for _, g in GRADES[:i]]
    open_keys = ' + '.join([f'(E{k + 2}-F{k + 2})' for k in range(i + 1)])
    lower_open = ' + '.join([f'COUNTIFS({G},"{g}",{E},"<>●",{I},"<>да")-COUNTIFS({G},"{g}",{E},"<>●",{F},"~*",{I},"вычеркнут")'
                             for g in lower]) or '0'
    ts.append([
        name,
        f'=COUNTIFS({G},"{name}")-COUNTIFS({G},"{name}",{F},"~*",{I},"вычеркнут")',
        f'=COUNTIFS({G},"{name}",{I},"да")',
        f'=INT((4*B{r}+4)/5)',
        f'=COUNTIFS({G},"{name}",{E},"●")',
        f'=COUNTIFS({G},"{name}",{E},"●",{I},"да")',
        f'={open_keys}',
        f'={lower_open}',
        f'=IF(AND(B{r}>0,G{r}=0,C{r}>=D{r},H{r}<=2),"да","нет")',
    ])
ts.append([])
# COUNTIF, not COUNTA: the empty cells of column I are written as empty strings
ts.append(['Грейд по матрице',
           f'=IF(COUNTIF({I},"да")+COUNTIF({I},"нет")=0,"заполните столбец «Итог после 1:1» на листе «Матрица»",'
           'IF(I4="да","Senior",IF(I3="да","Middle",IF(I2="да","Junior","не засчитан: уровень стажёра или ассистента"))))'])
ts['A6'].font = bold
ts['B6'].font = bold
ts.append([])
ts.append(['Правило: засчитан, когда закрыты все ключевые пункты (●) этого грейда и грейдов ниже, закрыто не меньше '
           '80% пунктов грейда (без вычеркнутых, с округлением вверх) и в грейдах ниже открыты не больше двух обычных '
           'пунктов. Пункт закрыт, когда в столбце «Итог после 1:1» стоит «да».'])
ts['A8'].alignment = wrap
ts.merge_cells('A8:I8')
ts.row_dimensions[8].height = 48
for col, w in zip('ABCDEFGHI', [18, 14, 10, 12, 10, 12, 16, 14, 11]):
    ts.column_dimensions[col].width = w
for row in ts.iter_rows(min_row=2, max_row=4):
    for c in row:
        c.border = Border(bottom=hair)
landscape(ts)

# --- Ключевые пункты ---------------------------------------------------------
ks.append(['Грейд', '№', 'Пункт', 'Что попросить показать'])
header(ks)
keys = [s for cat in d['categories'] for s in cat['skills'] if s.get('key')]
keys.sort(key=lambda s: ORDER.index(s['grade']))  # stable: the matrix order within a grade
for s in keys:
    ks.append([NAME[s['grade']], s['id'], s['short'], s['evidence']])
for row in ks.iter_rows(min_row=2, max_row=ks.max_row):
    for c in row:
        c.alignment = wrap
        c.border = Border(bottom=hair)
for col, w in zip('ABCD', [9, 6, 45, 70]):
    ks.column_dimensions[col].width = w
ks.freeze_panes = 'A2'
landscape(ks, '1:1')

# --- Как пользоваться --------------------------------------------------------
hs.column_dimensions['A'].width = 110
fill = [
    'Лист «Матрица», по строке на навык:',
    '   — «Методист»: отметки методиста, «да» или «нет» (шаг 2);',
    '   — «Руководитель»: ваши отметки (шаг 3);',
    '   — «Итог после 1:1»: итоговая отметка после встречи (шаг 4); у пунктов со звёздочкой (*) '
    'здесь можно выбрать «вычеркнут» (шаг 1);',
    '   — «Пример или комментарий»: пример, который показал методист, или решение по пункту.',
    'Лист «Итог» считает грейд по столбцу «Итог после 1:1» (шаг 5).',
    'Самопроверка: ставьте свои отметки сразу в столбец «Итог после 1:1».',
    'Ключевые пункты и что попросить показать — на листе «Ключевые пункты».',
]
lines = [(d['title'], 'title'), (d['lead'], None), ('', None), ('Как заполнять таблицу', 'h')]
lines += [(x, None) for x in fill]
lines += [('', None), (d['howHeading'], 'h')]
for k, s in enumerate(d['steps'], 1):
    lines.append((step_line(k, s), None))
    for x in s.get('items') or []:
        lines.append((f'   — {x}', None))
    if s.get('after'):
        lines.append((f"   {s['after']}", None))
if d.get('howAfter'):
    lines.append(('', None))
    lines += [(x, None) for x in d['howAfter']]
lines += [('', None), ('Обозначения', 'h'), (f"● — {d['legend']['key']}", None), (f"* — {d['legend']['star']}", None),
          ('', None), (f'Страница с матрицей и подсчётом грейда: {page_url}', page_url),
          (f"{d.get('version') or ''} Даниил Евстратов, {SITE}".strip(), SITE + '/')]
for text, kind in lines:
    hs.append([text])
    c = hs.cell(row=hs.max_row, column=1)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if kind == 'title':
        c.font = Font(bold=True, size=14)
    elif kind == 'h':
        c.font = bold
    elif kind:
        c.hyperlink = f'{kind}?{UTM}'
        c.font = Font(underline='single')
landscape(hs)

wb.properties.title = d['title']
wb.properties.creator = 'Даниил Евстратов'
wb.active = 0
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f'wrote {out.relative_to(ROOT)}: {n} skills, {len(keys)} key items, {out.stat().st_size // 1024} KB')
